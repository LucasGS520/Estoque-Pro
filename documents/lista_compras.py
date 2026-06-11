""" Geração do documento Word da Lista Final de Compras.

Recebe o DataFrame exibido na tela de Compras e cria um `.docx` agrupado por
categoria, usando os estilos base definidos em `documents.base_doc`.
"""

from pathlib import Path

import pandas as pd
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

_COR_PRETO = RGBColor(0x1A, 0x1A, 0x1A)

from documents.base_doc import criar_documento, adicionar_tabela, salvar_com_timestamp

_TITULO = "Lista Final de Itens e Quantidade para Compra"


def gerar_lista_compras(
    df: pd.DataFrame,
    periodo: str = "",
    data_geracao: str = "",
) -> Path:
    """
    Gera a lista final de compras a partir de um DataFrame.

    Espera colunas: Item, Categoria, Tipo, Estoque Atual, Qtd Compra.
    Retorna o caminho do arquivo salvo.
    """
    doc = criar_documento(_TITULO)

    if periodo:
        sub_p = doc.add_paragraph()
        lbl = sub_p.add_run("Período de referência:  ")
        lbl.bold = True
        lbl.font.size = Pt(9)
        per = sub_p.add_run(periodo)
        per.font.size = Pt(9)
        per.italic = True

    doc.add_paragraph()

    categorias = sorted(df["Categoria"].unique()) if "Categoria" in df.columns else ["—"]

    for cat in categorias:
        df_cat = df[df["Categoria"] == cat] if "Categoria" in df.columns else df
        if df_cat.empty:
            continue

        # "—" é a categoria dos itens de Compra Avulsa
        heading_label = "Compra Avulsa" if cat == "—" else cat
        heading = doc.add_heading(heading_label, level=3)
        heading.runs[0].font.size = Pt(11)
        heading.runs[0].font.color.rgb = _COR_PRETO

        headers = ["Item", "Estoque Atual", "Qtd Compra", "Tipo"]
        rows = []
        for _, row in df_cat.iterrows():
            atual = row.get("Estoque Atual")
            rows.append([
                str(row.get("Item", "")),
                str(int(atual)) if pd.notna(atual) else "—",
                str(row.get("Qtd Compra", "")),
                str(row.get("Tipo", "")),
            ])

        adicionar_tabela(doc, headers, rows)
        doc.add_paragraph()

    # Rodapé — data já está no topo; aqui só o total
    n_itens = len(df)
    doc.add_paragraph()
    rodape_p = doc.add_paragraph()
    rodape_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = rodape_p.add_run(f"{n_itens} item(ns) para compra")
    run.font.size = Pt(8)
    run.italic = True

    return salvar_com_timestamp(doc, "lista_compras")
