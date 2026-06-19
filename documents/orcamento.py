"""Geração do Documento de Cotação para Aprovação.

Recebe um objeto Orcamento (ORM) com itens e cotações já carregados e gera
um .docx estruturado para decisão do diretor, com cabeçalho de identificação,
resumo financeiro, tabela principal por categoria e cotações alternativas.
"""

from itertools import groupby
from pathlib import Path

from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core import tempo
from data.models import Orcamento
from documents.base_doc import (
    adicionar_tabela,
    criar_documento,
    salvar_com_timestamp,
    _set_cell_bg,
)

_COR_PRETO  = RGBColor(0x1A, 0x1A, 0x1A)
_COR_VERDE  = RGBColor(0x16, 0xA3, 0x4A)
_COR_BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
_COR_CINZA  = RGBColor(0x64, 0x74, 0x8B)
_HEX_CAT    = "CC0000"

_STATUS_PT: dict[str, str] = {
    "RASCUNHO":               "Rascunho",
    "EM_COTACAO":             "Em cotação",
    "PRONTO_PARA_APROVACAO":  "Pronto para aprovação",
    "ENVIADO_APROVACAO":      "Aguardando aprovação",
    "APROVADO":               "Aprovado",
    "APROVADO_PARCIAL":       "Aprovado parcialmente",
    "DEVOLVIDO_REVISAO":      "Devolvido para revisão",
    "REPROVADO":              "Reprovado",
    "CANCELADO":              "Cancelado",
}


def _fmt(value: float | None) -> str:
    if value is None:
        return "—"
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _periodo_fmt(periodo) -> str:
    meses = [
        "", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    ]
    return f"{meses[periodo.month]}/{periodo.year}"


def gerar_documento_cotacao(orc: Orcamento) -> Path:
    """Gera o Documento de Cotação para Aprovação a partir de um Orcamento ORM.

    Os itens devem estar carregados com suas cotações (eager load via joinedload).
    Retorna o caminho do arquivo gerado em EXPORTS_DIR.
    """
    doc = criar_documento("Documento de Cotação para Aprovação")

    # ── Identificação do orçamento ──────────────────────────────────────────
    periodo_str = _periodo_fmt(orc.periodo_referencia)
    status_str  = _STATUS_PT.get(orc.status, orc.status)

    info_dados = [
        ("Período:",         periodo_str),
        ("Orçamento nº:",    orc.numero),
        ("Versão:",          str(orc.versao)),
        ("Data:",            tempo.agora().strftime("%d/%m/%Y")),
        ("Responsável:",     orc.responsavel or "—"),
        ("Status:",          status_str),
    ]

    tbl_info = doc.add_table(rows=len(info_dados), cols=2)
    tbl_info.style = "Table Grid"
    for i, (lbl, val) in enumerate(info_dados):
        c = tbl_info.rows[i].cells
        c[0].text = lbl
        r0 = c[0].paragraphs[0].runs[0]
        r0.bold = True
        r0.font.size = Pt(9)
        c[1].text = val
        r1 = c[1].paragraphs[0].runs[0]
        r1.font.size = Pt(9)
        for cell in c:
            _set_cell_bg(cell, "F9FAFB")

    doc.add_paragraph()

    # ── Resumo financeiro (baseado no menor preço por item) ────────────────
    itens_sem_cotacao = [i for i in orc.itens if not i.cotacoes]

    total_bruto = 0.0
    frete_total = 0.0
    total_geral = 0.0
    origens: set[str] = set()

    for item in orc.itens:
        if not item.cotacoes:
            continue
        cot = min(item.cotacoes, key=lambda c: c.total_calculado)
        total_bruto += item.quantidade_orcada * cot.preco_unitario - cot.desconto + cot.outros_custos
        frete_total += cot.frete
        total_geral += cot.total_calculado
        origens.add(cot.origem)

    res_heading = doc.add_heading("Resumo Financeiro", level=3)
    for run in res_heading.runs:
        run.font.color.rgb = _COR_PRETO

    resumo_dados = [
        ("Itens orçados:",         str(len(orc.itens))),
        ("Origens/fornecedores:",  str(len(origens)) if origens else "—"),
        ("Total bruto:",           _fmt(round(total_bruto, 2))),
        ("Frete total:",           _fmt(round(frete_total, 2))),
        ("Total geral:",           _fmt(round(total_geral, 2))),
    ]

    tbl_res = doc.add_table(rows=len(resumo_dados), cols=2)
    tbl_res.style = "Table Grid"
    for i, (lbl, val) in enumerate(resumo_dados):
        c = tbl_res.rows[i].cells
        c[0].text = lbl
        r0 = c[0].paragraphs[0].runs[0]
        r0.bold = True
        r0.font.size = Pt(9)
        c[1].text = val
        r1 = c[1].paragraphs[0].runs[0]
        r1.font.size = Pt(9)
        if lbl == "Total geral:":
            r1.bold = True
        for cell in c:
            _set_cell_bg(cell, "F9FAFB")

    doc.add_paragraph()

    # ── Cotações por item agrupadas por categoria ───────────────────────────
    itens_ord = sorted(orc.itens, key=lambda i: (i.categoria_nome or "", i.nome_item))

    main_heading = doc.add_heading("Cotações por Item", level=3)
    for run in main_heading.runs:
        run.font.color.rgb = _COR_PRETO

    cot_headers = ["Fornecedor / Origem", "Unit.", "Frete", "Desconto", "Total", "Data"]
    cot_widths  = [4.5, 1.8, 1.6, 1.6, 1.8, 1.8]

    for cat, grupo in groupby(itens_ord, key=lambda i: i.categoria_nome or "—"):
        itens_cat = list(grupo)

        # Cabeçalho de categoria
        cat_tbl = doc.add_table(rows=1, cols=1)
        cat_tbl.style = "Table Grid"
        cat_cell = cat_tbl.rows[0].cells[0]
        cat_cell.text = cat.upper()
        run_ct = cat_cell.paragraphs[0].runs[0]
        run_ct.bold = True
        run_ct.font.size = Pt(9)
        run_ct.font.color.rgb = _COR_BRANCO
        _set_cell_bg(cat_cell, _HEX_CAT)

        for item in itens_cat:
            # Linha de identificação do item
            item_p = doc.add_paragraph()
            item_run = item_p.add_run(
                f"{item.nome_item}  —  Qtd: {item.quantidade_orcada}"
                f"  ({'Avulsa' if item.tipo == 'AVULSA' else 'Estoque'})"
            )
            item_run.bold = True
            item_run.font.size = Pt(9)

            if not item.cotacoes:
                sc_p = doc.add_paragraph()
                sc_p.paragraph_format.left_indent = Cm(0.5)
                sc_p.add_run("Sem cotações registradas.").font.size = Pt(8)
                continue

            cotacoes_ord = sorted(item.cotacoes, key=lambda c: c.total_calculado)
            menor_total = cotacoes_ord[0].total_calculado

            # Tabela comparativa de cotações
            tbl = doc.add_table(rows=1, cols=len(cot_headers))
            tbl.style = "Table Grid"
            hdr_cells = tbl.rows[0].cells
            for j, h in enumerate(cot_headers):
                hdr_cells[j].text = h
                r = hdr_cells[j].paragraphs[0].runs[0]
                r.bold = True
                r.font.size = Pt(8)
                r.font.color.rgb = _COR_BRANCO
                _set_cell_bg(hdr_cells[j], "1A1A1A")

            for idx, cot in enumerate(cotacoes_ord):
                bg = "F0FDF4" if cot.total_calculado == menor_total else ("F5F5F5" if idx % 2 == 0 else "FFFFFF")
                row_cells = tbl.add_row().cells
                prefixo = "★ " if cot.total_calculado == menor_total else ""
                valores = [
                    prefixo + cot.origem,
                    _fmt(cot.preco_unitario),
                    _fmt(cot.frete) if cot.frete else "—",
                    _fmt(cot.desconto) if cot.desconto else "—",
                    _fmt(cot.total_calculado),
                    cot.data_cotacao.strftime("%d/%m/%Y") if cot.data_cotacao else "—",
                ]
                for j, val in enumerate(valores):
                    row_cells[j].text = val
                    run_v = row_cells[j].paragraphs[0].runs[0]
                    run_v.font.size = Pt(8)
                    if cot.total_calculado == menor_total:
                        run_v.bold = True
                        run_v.font.color.rgb = _COR_VERDE
                    _set_cell_bg(row_cells[j], bg)
                    if j < len(cot_widths):
                        row_cells[j].width = Cm(cot_widths[j])

            doc.add_paragraph()

        doc.add_paragraph()

    # ── Itens sem cotação ───────────────────────────────────────────────────
    if itens_sem_cotacao:
        sc_heading = doc.add_heading("Itens Sem Cotação", level=3)
        for run in sc_heading.runs:
            run.font.color.rgb = _COR_PRETO
        adicionar_tabela(
            doc,
            headers=["Item", "Categoria", "Qtd", "Tipo"],
            rows=[
                (i.nome_item, i.categoria_nome, i.quantidade_orcada, i.tipo)
                for i in itens_sem_cotacao
            ],
        )
        doc.add_paragraph()

    # ── Campo de decisão ────────────────────────────────────────────────────
    dec_heading = doc.add_heading("Decisão de Aprovação", level=3)
    for run in dec_heading.runs:
        run.font.color.rgb = _COR_PRETO

    if orc.decidido_em:
        decisao_dados = [
            ("Aprovado por:",   orc.responsavel or "—"),
            ("Data:",           orc.decidido_em.strftime("%d/%m/%Y %H:%M")),
            ("Status:",         status_str),
            ("Observação:",     orc.observacao_diretor or "—"),
        ]
    else:
        decisao_dados = [
            ("Aprovado por:",  ""),
            ("Data:",          ""),
            ("Decisão:",       "[ ] Aprovado  [ ] Aprovado parcialmente  [ ] Revisão  [ ] Reprovado"),
            ("Observação:",    ""),
        ]

    tbl_dec = doc.add_table(rows=len(decisao_dados), cols=2)
    tbl_dec.style = "Table Grid"
    for i, (lbl, val) in enumerate(decisao_dados):
        c = tbl_dec.rows[i].cells
        c[0].text = lbl
        r0 = c[0].paragraphs[0].runs[0]
        r0.bold = True
        r0.font.size = Pt(9)
        c[1].text = val
        if c[1].paragraphs[0].runs:
            c[1].paragraphs[0].runs[0].font.size = Pt(9)

    return salvar_com_timestamp(doc, f"orcamento_{orc.numero}_v{orc.versao}")
